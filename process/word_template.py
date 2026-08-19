import re

from docx import Document



def get_placeholders (document_path: str) -> set[str]:

    document = Document(document_path)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    pattern = r"\{\{(.*?)\}\}"

    placeholders = {
        match.strip().upper()
        for match in re.findall(pattern, text)
    }

    return placeholders